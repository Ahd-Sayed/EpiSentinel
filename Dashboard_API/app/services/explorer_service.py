from sqlalchemy import func, or_, and_, desc, asc
from app.models.alerts import WeeklyAlertFact
from app.models.lstm import LSTMPredictionFact
from app.models.prediction_detail import PredictionDetail
from app.models.signals import WeeklyVisitsAndSymptomsFact
from app.models.patients import PatientVisitFact
from app.models.geography import GeographyDim
from app.models.disease import DiseaseDim
from app.extensions import db
import io
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO
from datetime import datetime

class ExplorerService:
    @staticmethod
    def get_filters():
        alerts_cities = db.session.query(WeeklyAlertFact.city.distinct()).order_by(WeeklyAlertFact.city).all()
        alerts_diseases = db.session.query(WeeklyAlertFact.disease.distinct()).order_by(WeeklyAlertFact.disease).all()
        
        lstm_countries = db.session.query(LSTMPredictionFact.country.distinct()).order_by(LSTMPredictionFact.country).all()
        lstm_cities = db.session.query(LSTMPredictionFact.city.distinct()).order_by(LSTMPredictionFact.city).all()
        
        patient_cities = db.session.query(PatientVisitFact.city.distinct()).order_by(PatientVisitFact.city).all()
        patient_diseases = db.session.query(PatientVisitFact.disease.distinct()).order_by(PatientVisitFact.disease).all()
        patient_seasons = db.session.query(PatientVisitFact.season.distinct()).order_by(PatientVisitFact.season).all()
        
        return {
            "alerts_cities":    [c[0] for c in alerts_cities if c[0]],
            "alerts_diseases":  [d[0] for d in alerts_diseases if d[0]],
            "lstm_countries":   [c[0] for c in lstm_countries if c[0]],
            "lstm_cities":      [c[0] for c in lstm_cities if c[0]],
            "patient_cities":   [c[0] for c in patient_cities if c[0]],
            "patient_diseases": [d[0] for d in patient_diseases if d[0]],
            "patient_seasons":  [s[0] for s in patient_seasons if s[0]],
        }

    @staticmethod
    def _apply_sorting(query, model, sort_by, sort_dir):
        if not sort_by:
            return query
        col_attr = getattr(model, sort_by.lower(), None)
        if col_attr:
            if sort_dir == "desc":
                query = query.order_by(desc(col_attr))
            else:
                query = query.order_by(asc(col_attr))
        return query

    # ── Tab 1: Weekly Alerts ──────────────────────────────────────────────────
    @staticmethod
    def get_alerts(cities=None, diseases=None, levels=None, date_from=None, date_to=None,
                   score_min=None, search=None, sort_by=None, sort_dir='asc', page=1, per_page=50, prediction_id=None):
        
        query = WeeklyAlertFact.query
        if prediction_id:
            query = query.filter(WeeklyAlertFact.prediction_id == prediction_id)
        else:
            query = query.filter(WeeklyAlertFact.prediction_id.is_(None))

        
        if cities:
            query = query.filter(WeeklyAlertFact.city.in_(cities))
        if diseases:
            query = query.filter(WeeklyAlertFact.disease.in_(diseases))
        if levels:
            query = query.filter(WeeklyAlertFact.alert_level.in_(levels))
        if date_from:
            try:
                dt = datetime.strptime(date_from, "%Y-%m-%d").date()
                query = query.filter(WeeklyAlertFact.week_start >= dt)
            except ValueError:
                pass
        if date_to:
            try:
                dt = datetime.strptime(date_to, "%Y-%m-%d").date()
                query = query.filter(WeeklyAlertFact.week_start <= dt)
            except ValueError:
                pass
        if score_min is not None and score_min != "":
            query = query.filter(WeeklyAlertFact.fusion_score >= int(score_min))
            
        if search:
            query = query.filter(
                or_(
                    WeeklyAlertFact.city.like(f"%{search}%"),
                    WeeklyAlertFact.disease.like(f"%{search}%"),
                    WeeklyAlertFact.alert_level.like(f"%{search}%")
                )
            )
            
        query = ExplorerService._apply_sorting(query, WeeklyAlertFact, sort_by or "week_start", sort_dir or "desc")
        
        pagination = query.paginate(page=page, per_page=per_page, error_out=False)
        
        records = []
        for a in pagination.items:
            records.append({
                "CITY": a.city,
                "DISEASE": a.disease,
                "WEEK_START": str(a.week_start),
                "IF_FLAG": a.if_flag,
                "DBSCAN_FLAG": a.dbscan_flag,
                "lstm_pred": a.lstm_pred,
                "lstm_prob": round(float(a.lstm_prob), 4) if a.lstm_prob is not None else 0.0,
                "fusion_score": a.fusion_score,
                "alert_level": a.alert_level,
                "total_visits": a.total_visits or 0.0
            })
            
        return records, pagination.total, pagination.page, pagination.pages

    # ── Tab 2: LSTM Predictions ───────────────────────────────────────────────
    @staticmethod
    def get_lstm(country=None, cities=None, diseases=None, date_from=None, date_to=None,
                 search=None, sort_by=None, sort_dir='asc', page=1, per_page=50, prediction_id=None):
        
        if prediction_id:
            query = PredictionDetail.query.filter_by(prediction_id=prediction_id)
            if country:
                query = query.filter(PredictionDetail.country == country)
            if cities:
                query = query.filter(PredictionDetail.city.in_(cities))
            if diseases:
                query = query.filter(PredictionDetail.disease.in_(diseases))
            if date_from:
                query = query.filter(PredictionDetail.week >= date_from)
            if date_to:
                query = query.filter(PredictionDetail.week <= date_to)
            if search:
                query = query.filter(or_(
                    PredictionDetail.country.like(f"%{search}%"),
                    PredictionDetail.city.like(f"%{search}%"),
                    PredictionDetail.disease.like(f"%{search}%")
                ))
                
            total = query.count()
            pgs = (total + per_page - 1) // per_page
            items = query.order_by(PredictionDetail.id.desc()).offset((page - 1) * per_page).limit(per_page).all()
            
            result_data = []
            for item in items:
                result_data.append({
                    "lstm_id": item.id,
                    "COUNTRY": item.country,
                    "CITY": item.city,
                    "DISEASE": item.disease,
                    "target_date": str(item.week),
                    "lstm_prob": round(item.lstm_probability, 3) if item.lstm_probability else 0.0,
                    "lstm_pred": item.lstm_flag,
                    "y_true": 0
                })
            return result_data, total, page, pgs

        query = LSTMPredictionFact.query

        
        if country:
            query = query.filter_by(country=country)
        if cities:
            query = query.filter(LSTMPredictionFact.city.in_(cities))
        if diseases:
            query = query.filter(LSTMPredictionFact.disease.in_(diseases))
        if date_from:
            try:
                dt = datetime.strptime(date_from, "%Y-%m-%d").date()
                query = query.filter(LSTMPredictionFact.target_date >= dt)
            except ValueError:
                pass
        if date_to:
            try:
                dt = datetime.strptime(date_to, "%Y-%m-%d").date()
                query = query.filter(LSTMPredictionFact.target_date <= dt)
            except ValueError:
                pass
                
        if search:
            query = query.filter(
                or_(
                    LSTMPredictionFact.country.like(f"%{search}%"),
                    LSTMPredictionFact.city.like(f"%{search}%"),
                    LSTMPredictionFact.disease.like(f"%{search}%")
                )
            )
            
        query = ExplorerService._apply_sorting(query, LSTMPredictionFact, sort_by or "target_date", sort_dir or "desc")
        
        pagination = query.paginate(page=page, per_page=per_page, error_out=False)
        
        records = []
        for a in pagination.items:
            records.append({
                "COUNTRY": a.country,
                "CITY": a.city,
                "DISEASE": a.disease,
                "target_date": str(a.target_date),
                "lstm_prob": round(float(a.lstm_prob), 4),
                "lstm_pred": a.lstm_pred,
                "y_true": a.y_true
            })
            
        return records, pagination.total, pagination.page, pagination.pages

    # ── Tab 3: Merged Signals ─────────────────────────────────────────────────
    @staticmethod
    def get_signals(cities=None, diseases=None, date_from=None, date_to=None,
                    search=None, sort_by=None, sort_dir="asc", page=1, per_page=50, prediction_id=None):
        
        query = WeeklyVisitsAndSymptomsFact.query
        if prediction_id:
            query = query.filter(WeeklyVisitsAndSymptomsFact.prediction_id == prediction_id)
        else:
            query = query.filter(WeeklyVisitsAndSymptomsFact.prediction_id.is_(None)).filter_by(is_overlap=True) # overlap tab matches explorer
        
        if cities:
            query = query.filter(WeeklyVisitsAndSymptomsFact.city.in_(cities))
        if diseases:
            query = query.filter(WeeklyVisitsAndSymptomsFact.disease.in_(diseases))
        if date_from:
            try:
                dt = datetime.strptime(date_from, "%Y-%m-%d").date()
                query = query.filter(WeeklyVisitsAndSymptomsFact.week_start >= dt)
            except ValueError:
                pass
        if date_to:
            try:
                dt = datetime.strptime(date_to, "%Y-%m-%d").date()
                query = query.filter(WeeklyVisitsAndSymptomsFact.week_start <= dt)
            except ValueError:
                pass
                
        if search:
            query = query.filter(
                or_(
                    WeeklyVisitsAndSymptomsFact.city.like(f"%{search}%"),
                    WeeklyVisitsAndSymptomsFact.disease.like(f"%{search}%")
                )
            )
            
        query = ExplorerService._apply_sorting(query, WeeklyVisitsAndSymptomsFact, sort_by or "week_start", sort_dir or "desc")
        
        pagination = query.paginate(page=page, per_page=per_page, error_out=False)
        
        records = []
        for a in pagination.items:
            records.append({
                "CITY": a.city,
                "DISEASE": a.disease,
                "WEEK_START": str(a.week_start),
                "IF_FLAG": a.if_flag,
                "DBSCAN_FLAG": a.dbscan_flag,
                "lstm_prob": round(float(a.lstm_prob), 4) if a.lstm_prob is not None else 0.0,
                "lstm_pred": a.lstm_pred,
                "true_outbreak": a.true_outbreak or 0,
                "total_visits": a.total_visits or 0.0
            })
            
        return records, pagination.total, pagination.page, pagination.pages

    # ── Tab 4: Patient Records (sample / fully indexed) ───────────────────────
    @staticmethod
    def get_patients(cities=None, diseases=None, gender=None, season=None, is_outbreak=None,
                     if_pred=None, dbscan_pred=None, search=None, sort_by=None, sort_dir="asc", page=1, per_page=50, prediction_id=None):
        
        query = PatientVisitFact.query
        if prediction_id:
            query = query.filter(PatientVisitFact.prediction_id == prediction_id)
        else:
            query = query.filter(PatientVisitFact.prediction_id.is_(None))

        
        if cities:
            query = query.filter(PatientVisitFact.city.in_(cities))
        if diseases:
            query = query.filter(PatientVisitFact.disease.in_(diseases))
        if gender:
            query = query.filter_by(gender=gender)
        if season:
            query = query.filter_by(season=season)
        if is_outbreak is not None and is_outbreak != "":
            query = query.filter_by(is_outbreak=int(is_outbreak))
        if if_pred is not None and if_pred != "":
            query = query.filter_by(if_pred=int(if_pred))
        if dbscan_pred is not None and dbscan_pred != "":
            query = query.filter_by(dbscan_pred=int(dbscan_pred))
            
        if search:
            query = query.filter(
                or_(
                    PatientVisitFact.city.like(f"%{search}%"),
                    PatientVisitFact.disease.like(f"%{search}%"),
                    PatientVisitFact.gender.like(f"%{search}%"),
                    PatientVisitFact.season.like(f"%{search}%"),
                    PatientVisitFact.hospital_name.like(f"%{search}%")
                )
            )
            
        # Ensure we map frontend capitalized sort_by cols (e.g. VISIT_DATE) to SQLAlchemy lower attributes
        sort_col = sort_by or "visit_date"
        if sort_col.upper() == "VISIT_DATE":
            sort_col = "visit_date"
        elif sort_col.upper() == "AGE_AT_VISIT":
            sort_col = "age_at_visit"
        elif sort_col.upper() == "COMORBIDITY_COUNT":
            sort_col = "comorbidity_count"
        elif sort_col.upper() == "IS_OUTBREAK":
            sort_col = "is_outbreak"
            
        query = ExplorerService._apply_sorting(query, PatientVisitFact, sort_col, sort_dir or "desc")
        
        pagination = query.paginate(page=page, per_page=per_page, error_out=False)
        
        records = []
        for a in pagination.items:
            records.append({
                "VISIT_DATE": a.visit_date.strftime("%Y-%m-%d"),
                "CITY": a.city,
                "DISEASE": a.disease,
                "AGE_AT_VISIT": round(float(a.age_at_visit), 1) if a.age_at_visit is not None else 0.0,
                "GENDER": a.gender,
                "SEASON": a.season,
                "COMORBIDITY_COUNT": a.comorbidity_count,
                "IS_OUTBREAK": a.is_outbreak,
                "if_score": round(float(a.if_score), 4) if a.if_score is not None else 0.0,
                "if_pred": a.if_pred,
                "dbscan_score": round(float(a.dbscan_score), 4) if a.dbscan_score is not None else 0.0,
                "dbscan_pred": a.dbscan_pred
            })
            
        return records, pagination.total, pagination.page, pagination.pages

    # ── Excel and CSV exporter ────────────────────────────────────────────────
    @staticmethod
    def get_export_data(source, limit=5000):
        # Queries up to 'limit' rows
        if source == "alerts":
            query = WeeklyAlertFact.query.order_by(WeeklyAlertFact.week_start.desc())
            cols = ["city", "disease", "week_start", "target_date", "if_flag", "dbscan_flag", 
                    "lstm_pred", "lstm_prob", "fusion_score", "alert_level", "total_visits"]
            data = query.limit(limit).all()
            return cols, [{c.upper(): getattr(a, c.lower()) for c in cols} for a in data]
            
        elif source == "lstm":
            query = LSTMPredictionFact.query.order_by(LSTMPredictionFact.target_date.desc())
            cols = ["country", "city", "disease", "target_date", "lstm_prob", "lstm_pred", "y_true"]
            data = query.limit(limit).all()
            return cols, [{c.upper() if c != "target_date" else c: getattr(a, c) for c in cols} for a in data]
            
        elif source == "signals":
            query = WeeklyVisitsAndSymptomsFact.query.filter_by(is_overlap=True).order_by(WeeklyVisitsAndSymptomsFact.week_start.desc())
            cols = ["city", "disease", "week_start", "if_flag", "dbscan_flag", "lstm_prob", "lstm_pred", "true_outbreak", "total_visits"]
            data = query.limit(limit).all()
            return cols, [{c.upper(): getattr(a, c.lower()) for c in cols} for a in data]
            
        elif source == "patients":
            query = PatientVisitFact.query.order_by(PatientVisitFact.visit_date.desc())
            cols = ["visit_date", "city", "disease", "age_at_visit", "gender", "season", "comorbidity_count", "is_outbreak"]
            data = query.limit(limit).all()
            
            records = []
            for a in data:
                records.append({
                    "VISIT_DATE": a.visit_date.strftime("%Y-%m-%d"),
                    "CITY": a.city,
                    "DISEASE": a.disease,
                    "AGE_AT_VISIT": round(float(a.age_at_visit), 1) if a.age_at_visit is not None else 0.0,
                    "GENDER": a.gender,
                    "SEASON": a.season,
                    "COMORBIDITY_COUNT": a.comorbidity_count,
                    "IS_OUTBREAK": a.is_outbreak
                })
            return [c.upper() for c in cols], records
            
        return [], []

    @staticmethod
    def generate_excel_bytes(cols, records):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EpiGuard Export"
        
        # Style rules
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        cell_font = Font(name="Segoe UI", size=10)
        
        # Headers
        ws.append([c.upper() for c in cols])
        for col_num in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            
        for r in records:
            ws.append([r.get(c.upper() if c != "target_date" else c, "") for c in cols])
            
        # Format grid dimensions
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        out = BytesIO()
        wb.save(out)
        return out.getvalue()
